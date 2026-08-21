"""JSON and XLSX exporters.

Everything here is fed already-redacted data. JSON is dependency-free; XLSX uses
openpyxl. PDF is a Phase-3+ concern (weasyprint) and intentionally omitted to
avoid a heavy native dependency in the core.
"""

from __future__ import annotations

import io
import json
from collections import Counter

from app.execution.evidence import verify_chain
from app.schemas.execution import Execution
from app.schemas.finding import Finding
from app.schemas.testcase import TestCase


def export_json(
    issue_key: str,
    target: str,
    tests: list[TestCase],
    executions: list[Execution],
    findings: list[Finding],
    coverage: list[dict] | None = None,
    plan_review=None,
    run_assessment=None,
    derived_verdicts=None,
    report_manifest=None,
    plan_stale: bool = False,
    uncovered_endpoints: list[str] | None = None,
) -> str:
    """The machine-readable whole. The two agent artefacts are included under
    their own keys rather than merged into `executions`: an adjudication is an
    opinion about an execution, and folding it into the execution record would
    put an unhashed, model-derived field inside the object whose hash is supposed
    to make it tamper-evident.

    `plan_stale` says the test cases below were designed from an endpoint list
    that has since been edited — surfaced here because this export, unlike the
    live dashboard, is a snapshot a reader may act on without ever seeing the
    dashboard's own staleness banner.

    `uncovered_endpoints` names endpoint signatures a PoC-derived test actually
    hits that never made the endpoint list at all — a different, timing-
    independent gap (see `app.owasp.coverage.uncovered_poc_endpoints`)."""
    payload = {
        "issue_key": issue_key,
        "target": target,
        "summary": _summary(executions, findings),
        "coverage": coverage or [],
        "plan_stale": plan_stale,
        "uncovered_poc_endpoints": uncovered_endpoints or [],
        "test_cases": [t.model_dump(mode="json") for t in tests],
        "executions": [e.model_dump(mode="json") for e in executions],
        "findings": [f.model_dump(mode="json") for f in findings],
    }
    if plan_review is not None:
        payload["plan_review"] = plan_review.model_dump(mode="json")
    if run_assessment is not None:
        payload["run_assessment"] = run_assessment.model_dump(mode="json")
    if derived_verdicts is not None:
        payload["derived_verdicts"] = [item.model_dump(mode="json") for item in derived_verdicts]
    if report_manifest is not None:
        payload["report_manifest"] = report_manifest.model_dump(mode="json")
    return json.dumps(payload, indent=2, ensure_ascii=False)


def export_xlsx(
    issue_key: str,
    target: str,
    tests: list[TestCase],
    executions: list[Execution],
    findings: list[Finding],
    coverage: list[dict] | None = None,
    plan_stale: bool = False,
    uncovered_endpoints: list[str] | None = None,
) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    bold = Font(bold=True)
    warn = Font(bold=True, color="9C5700")
    uncovered_endpoints = uncovered_endpoints or []

    # Summary
    ws = wb.active
    ws.title = "Summary"
    summary = _summary(executions, findings)
    ws.append(["Security Assessment", issue_key])
    ws.append(["Target", target])
    if plan_stale:
        ws.append(["⚠ Test plan may be stale", "endpoints were edited after this plan was designed"])
        ws.cell(row=ws.max_row, column=1).font = warn
        ws.cell(row=ws.max_row, column=2).font = warn
    if uncovered_endpoints:
        ws.append(["⚠ Tests target endpoint(s) not in the endpoint list",
                   ", ".join(uncovered_endpoints)])
        ws.cell(row=ws.max_row, column=1).font = warn
        ws.cell(row=ws.max_row, column=2).font = warn
    ws.append([])
    ws.append(["Metric", "Count"])
    ws.cell(row=ws.max_row, column=1).font = ws.cell(row=ws.max_row, column=2).font = bold
    for k, v in summary["results"].items():
        ws.append([k, v])
    ws.append([])
    ws.append(["Severity", "Count"])
    for k, v in summary["severities"].items():
        ws.append([k, v])

    # Coverage
    ws = wb.create_sheet("Coverage")
    _header(ws, ["Category", "State", "Coverage %", "Existing PoC", "Generated"], bold)
    for r in coverage or []:
        ws.append([r.get("category"), r.get("state"), r.get("pct", 0),
                   r.get("existing_tests", 0), r.get("generated_tests", 0)])

    # Test cases
    suffix = " ".join(filter(None, ["(STALE)" if plan_stale else "",
                                    "(GAP)" if uncovered_endpoints else ""]))
    ws = wb.create_sheet(f"Test Cases {suffix}".rstrip() if suffix else "Test Cases")
    _header(ws, ["Test ID", "OWASP", "Severity", "Approval", "Destructive", "Mutation", "Title"], bold)
    for t in tests:
        ws.append([t.test_id, t.owasp_category.value, t.severity.value,
                   t.approval_status.value, t.is_destructive, t.attack_mutation.kind, t.title])

    # Executions
    ws = wb.create_sheet("Executions")
    _header(ws, ["Test ID", "OWASP", "Result", "Confidence", "Reason", "Evidence Hash"], bold)
    for e in executions:
        ws.append([e.test_id, e.owasp_category, e.verdict.result.value,
                   e.verdict.confidence.value, e.verdict.reason, e.evidence_hash[:16]])

    # Findings
    ws = wb.create_sheet("Findings")
    _header(ws, ["ID", "Title", "OWASP", "Severity", "Confidence", "Endpoint", "Recommendation"], bold)
    for f in findings:
        ws.append([f.finding_id, f.title, f.owasp_category.value, f.severity.value,
                   f.confidence.value, f.endpoint, f.recommendation])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _summary(executions: list[Execution], findings: list[Finding]) -> dict:
    results = Counter(e.verdict.result.value for e in executions)
    sev = Counter(f.severity.value for f in findings)
    return {
        "results": {k: results.get(k, 0) for k in
                    ["PASS", "FAIL", "INCONCLUSIVE", "BLOCKED", "ERROR", "SKIPPED", "TIMEOUT"]},
        "severities": {k: sev.get(k, 0) for k in
                       ["CRITICAL", "HIGH", "MEDIUM", "LOW", "INFO"]},
        "total_tests": len(executions),
        "total_findings": len(findings),
        "evidence_chain_ok": verify_chain(executions) if executions else None,
    }


def _header(ws, cols, bold) -> None:
    ws.append(cols)
    for cell in ws[1]:
        cell.font = bold
