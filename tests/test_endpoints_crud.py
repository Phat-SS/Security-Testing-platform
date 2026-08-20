"""Editing the attack surface.

The endpoint list is what the whole plan is derived from, and it is produced by
a regex over ticket prose: it misses endpoints written in a table, marks every
endpoint as requiring auth, and only finds object ids that appear in a path.
Until these routes existed the only way to correct any of that was to edit the
Jira ticket and import again, which threw away the plan and its approvals.
"""

import pytest
from starlette.testclient import TestClient


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path/'ep.db'}")
    monkeypatch.delenv("ENGAGEMENT_CONFIG", raising=False)
    from app.api.main import app

    with TestClient(app) as c:
        yield c


@pytest.fixture()
def repo():
    import app.api.main as main

    return lambda: main.state.repo


def _import(client) -> str:
    r = client.post("/import", data={"issue_key": "CRM-1234"}, follow_redirects=True)
    return r.url.path.rsplit("/", 1)[-1]


def _endpoints(repo, aid) -> list[dict]:
    return repo().get_assessment(aid).analysis_json["endpoints"]


def _signatures(repo, aid) -> list[str]:
    return [f"{ep['method']} {ep['path']}" for ep in _endpoints(repo, aid)]


# -- add --------------------------------------------------------------------


def test_add_endpoint_the_extractor_missed(client, repo):
    aid = _import(client)
    before = _signatures(repo, aid)

    r = client.post(f"/assessment/{aid}/endpoints", follow_redirects=True, data={
        "method": "post", "path": "/v2/orders", "auth_required": "true",
        "object_id_params": "orderId, customerId", "url_fields": "callbackUrl",
        "writes_properties": "true",
    })

    assert r.status_code == 200
    added = _endpoints(repo, aid)[-1]
    assert added["method"] == "POST"           # normalised
    assert added["path"] == "/v2/orders"
    assert added["object_id_params"] == ["orderId", "customerId"]
    assert added["url_fields"] == ["callbackUrl"]
    assert added["manual"] is True
    assert _signatures(repo, aid) == before + ["POST /v2/orders"]


def test_added_endpoint_is_designed_against(client, repo):
    """The point of adding one: it has to produce tests."""
    aid = _import(client)
    client.post(f"/assessment/{aid}/endpoints", follow_redirects=True,
                data={"method": "GET", "path": "/v2/invoices/{invoiceId}",
                      "auth_required": "true", "object_id_params": "invoiceId"})

    client.post(f"/assessment/{aid}/design", follow_redirects=True)

    paths = {t.request.path for t in repo().get_test_cases(aid)}
    assert "/v2/invoices/{invoiceId}" in paths


def test_a_duplicate_signature_is_refused_not_silently_merged(client, repo):
    aid = _import(client)
    existing = _signatures(repo, aid)[0]
    method, path = existing.split(" ", 1)

    r = client.post(f"/assessment/{aid}/endpoints", follow_redirects=True,
                    data={"method": method, "path": path})

    assert "already in the list" in r.text
    assert _signatures(repo, aid).count(existing) == 1


def test_a_path_without_a_leading_slash_is_refused(client, repo):
    """A test's path is joined onto the approved base URL — one carrying its own
    host would be a way around scope validation."""
    aid = _import(client)
    n = len(_endpoints(repo, aid))

    r = client.post(f"/assessment/{aid}/endpoints", follow_redirects=True,
                    data={"method": "GET", "path": "https://evil.example/x"})

    assert "must start with" in r.text
    assert len(_endpoints(repo, aid)) == n


# -- edit -------------------------------------------------------------------


def test_edit_keeps_the_row_in_place(client, repo):
    aid = _import(client)
    before = _signatures(repo, aid)
    assert len(before) >= 2, "fixture ticket should yield at least two endpoints"
    target = before[0]
    method, path = target.split(" ", 1)

    client.post(f"/assessment/{aid}/endpoints", follow_redirects=True, data={
        "replaces": target, "method": method, "path": path,
        "object_id_params": "customerId accountId",
    })

    after = _signatures(repo, aid)
    assert after == before, "an edit must not reorder the table"
    assert _endpoints(repo, aid)[0]["object_id_params"] == ["customerId", "accountId"]


def test_editing_can_turn_off_auth_required(client, repo):
    """The extractor marks every endpoint as authenticated, which generates API2
    probes against endpoints that have no credential to drop."""
    aid = _import(client)
    target = _signatures(repo, aid)[0]
    method, path = target.split(" ", 1)

    # an unchecked checkbox simply is not submitted
    client.post(f"/assessment/{aid}/endpoints", follow_redirects=True,
                data={"replaces": target, "method": method, "path": path})

    assert _endpoints(repo, aid)[0]["auth_required"] is False


def test_editing_can_mark_an_endpoint_expected_public(client, repo):
    """A separate switch from auth_required: declares a 200 without a credential
    is the correct, expected result for this endpoint, not a finding."""
    aid = _import(client)
    client.post(f"/assessment/{aid}/design", follow_redirects=True)
    target = _signatures(repo, aid)[0]
    method, path = target.split(" ", 1)

    client.post(f"/assessment/{aid}/endpoints", follow_redirects=True,
                data={"replaces": target, "method": method, "path": path,
                      "auth_required": "true", "expected_public": "true"})

    assert _endpoints(repo, aid)[0]["expected_public"] is True
    # Flipping it changes what the designer would generate, so the existing
    # plan (built before the flag was set) must show as stale.
    page = client.get(f"/assessment/{aid}").text
    assert "different endpoint list" in page

    client.post(f"/assessment/{aid}/design", follow_redirects=True)
    tests = repo().get_test_cases(aid)
    target_tests = [t for t in tests if t.request.path == path
                    and t.request.method == method.upper()
                    and t.owasp_category.value == "API2:2023"]
    assert target_tests
    assert all(t.attack_mutation.kind == "drop_auth" for t in target_tests)
    assert all(200 in t.expected.status_in for t in target_tests)


def test_editing_into_an_existing_signature_is_refused(client, repo):
    aid = _import(client)
    first, second = _signatures(repo, aid)[:2]
    method, path = second.split(" ", 1)

    r = client.post(f"/assessment/{aid}/endpoints", follow_redirects=True,
                    data={"replaces": first, "method": method, "path": path})

    assert "already in the list" in r.text
    assert _signatures(repo, aid)[:2] == [first, second]


# -- delete -----------------------------------------------------------------


def test_delete_endpoint(client, repo):
    aid = _import(client)
    target = _signatures(repo, aid)[0]

    r = client.post(f"/assessment/{aid}/endpoints/delete",
                    data={"signature": target}, follow_redirects=True)

    assert r.status_code == 200
    assert target not in _signatures(repo, aid)


def test_deleting_something_that_is_not_there_says_so(client, repo):
    aid = _import(client)

    r = client.post(f"/assessment/{aid}/endpoints/delete",
                    data={"signature": "GET /nope"}, follow_redirects=True)

    assert "Endpoint not found" in r.text


# -- OWASP mapping ----------------------------------------------------------


def _applicable(repo, aid) -> set[str]:
    return {m["category"] for m in repo().get_assessment(aid).analysis_json["owasp_mappings"]
            if m["applicability"] == "APPLICABLE"}


def test_adding_an_ssrf_surface_adds_the_ssrf_category(client, repo):
    aid = _import(client)

    client.post(f"/assessment/{aid}/endpoints", follow_redirects=True,
                data={"method": "POST", "path": "/v2/webhooks", "auth_required": "true",
                      "writes_properties": "true", "url_fields": "callbackUrl"})

    assert "API7:2023" in _applicable(repo, aid)


def test_an_endpoint_edit_never_downgrades_an_applicable_category(client, repo):
    """The mapping may have come from the Claude analyzer, which reads intent.
    Rebuilding it from the heuristic rules on every small edit would quietly
    replace that with something weaker, so an edit is additive-only and only an
    explicit re-analysis may remove a category."""
    aid = _import(client)
    before = _applicable(repo, aid)
    assert before

    for sig in list(_signatures(repo, aid)):
        client.post(f"/assessment/{aid}/endpoints/delete", data={"signature": sig})

    assert _endpoints(repo, aid) == []
    assert _applicable(repo, aid) == before


# -- re-analyze -------------------------------------------------------------


def test_reanalyze_rebuilds_the_list_and_keeps_hand_entered_rows(client, repo):
    aid = _import(client)
    extracted = _signatures(repo, aid)
    client.post(f"/assessment/{aid}/endpoints", follow_redirects=True,
                data={"method": "GET", "path": "/v2/hand-typed", "auth_required": "true"})
    # delete an extracted row so the rebuild has something to restore
    client.post(f"/assessment/{aid}/endpoints/delete", data={"signature": extracted[0]})

    r = client.post(f"/assessment/{aid}/reanalyze", follow_redirects=True)

    assert r.status_code == 200
    after = _signatures(repo, aid)
    assert extracted[0] in after, "re-analysis should restore what the ticket says"
    assert "GET /v2/hand-typed" in after, "a hand-entered endpoint must survive a rebuild"


def test_reanalyze_is_the_operation_that_may_drop_a_category(client, repo):
    aid = _import(client)
    client.post(f"/assessment/{aid}/endpoints", follow_redirects=True,
                data={"method": "POST", "path": "/v2/webhooks", "auth_required": "true",
                      "writes_properties": "true", "url_fields": "callbackUrl"})
    client.post(f"/assessment/{aid}/endpoints/delete", data={"signature": "POST /v2/webhooks"})
    assert "API7:2023" in _applicable(repo, aid)

    client.post(f"/assessment/{aid}/reanalyze", follow_redirects=True)

    assert "API7:2023" not in _applicable(repo, aid)
    # The mapping now reflects the ticket, not the accumulated history of edits:
    # it says exactly what a fresh import of the same issue would say.
    assert _applicable(repo, aid) == _applicable(repo, _import(client))


# -- stale plan -------------------------------------------------------------


def test_changing_endpoints_marks_an_existing_plan_stale(client, repo):
    aid = _import(client)
    client.post(f"/assessment/{aid}/design", follow_redirects=True)
    assert "different endpoint list" not in client.get(f"/assessment/{aid}").text

    client.post(f"/assessment/{aid}/endpoints", follow_redirects=True,
                data={"method": "GET", "path": "/v2/new-surface", "auth_required": "true"})

    page = client.get(f"/assessment/{aid}").text
    assert "different endpoint list" in page, (
        "a plan built for the old endpoint list must not look current"
    )
    assert "plan is stale" in page


def test_regenerating_clears_the_stale_warning(client):
    aid = _import(client)
    client.post(f"/assessment/{aid}/design", follow_redirects=True)
    client.post(f"/assessment/{aid}/endpoints", follow_redirects=True,
                data={"method": "GET", "path": "/v2/new-surface", "auth_required": "true"})

    client.post(f"/assessment/{aid}/design", follow_redirects=True)

    assert "different endpoint list" not in client.get(f"/assessment/{aid}").text


def test_a_filter_that_matches_nothing_does_not_hide_the_stale_warning(client):
    """The plan table is paged, so "does a plan exist" has to come from the
    plan-wide count and not from whatever the current page happens to hold."""
    aid = _import(client)
    client.post(f"/assessment/{aid}/design", follow_redirects=True)
    client.post(f"/assessment/{aid}/endpoints", follow_redirects=True,
                data={"method": "GET", "path": "/v2/new-surface", "auth_required": "true"})

    page = client.get(f"/assessment/{aid}?q=nothing-matches-this").text

    assert "Nothing matches this filter" in page, "the filter should have emptied the page"
    assert "different endpoint list" in page
    assert "Regenerate test plan" in page


# -- the stale warning must also travel with anything exported ---------------
#
# The dashboard's own banner (above) only reaches someone looking at the web
# UI. A tester who edits endpoints and then downloads a report, or posts the
# Jira comment, without revisiting the dashboard must not walk away with an
# artefact that silently claims a plan built for the old endpoint list is
# current.


def _stale_assessment(client) -> str:
    aid = _import(client)
    client.post(f"/assessment/{aid}/design", follow_redirects=True)
    client.post(f"/assessment/{aid}/endpoints", follow_redirects=True,
                data={"method": "GET", "path": "/v2/new-surface", "auth_required": "true"})
    return aid


def test_json_export_flags_a_stale_plan(client):
    import json

    import app.api.main as main

    aid = _stale_assessment(client)
    payload = json.loads(main.state.orch.export_json(aid))

    assert payload["plan_stale"] is True


def test_json_export_does_not_flag_a_fresh_plan(client):
    import json

    import app.api.main as main

    aid = _import(client)
    client.post(f"/assessment/{aid}/design", follow_redirects=True)
    payload = json.loads(main.state.orch.export_json(aid))

    assert payload["plan_stale"] is False


def test_xlsx_export_flags_a_stale_plan(client):
    import io

    from openpyxl import load_workbook

    import app.api.main as main

    aid = _stale_assessment(client)
    wb = load_workbook(io.BytesIO(main.state.orch.export_xlsx(aid)))

    assert "Test Cases (STALE)" in wb.sheetnames
    summary_cells = [str(c.value) for row in wb["Summary"].iter_rows() for c in row if c.value]
    assert any("stale" in v.lower() for v in summary_cells)


def test_xlsx_export_does_not_flag_a_fresh_plan(client):
    import io

    from openpyxl import load_workbook

    import app.api.main as main

    aid = _import(client)
    client.post(f"/assessment/{aid}/design", follow_redirects=True)
    wb = load_workbook(io.BytesIO(main.state.orch.export_xlsx(aid)))

    assert "Test Cases" in wb.sheetnames
    assert "Test Cases (STALE)" not in wb.sheetnames


def test_pdf_export_is_told_the_plan_is_stale(client, monkeypatch):
    import app.api.main as main
    import app.reporting.pdf as pdf_module

    captured = {}
    monkeypatch.setattr(pdf_module, "export_pdf",
                        lambda *a, **kw: captured.update(kw) or b"%PDF-fake")

    aid = _stale_assessment(client)
    main.state.orch.export_pdf(aid)

    assert captured.get("plan_stale") is True


def test_jira_comment_flags_a_stale_plan(client):
    import app.api.main as main

    aid = _stale_assessment(client)
    comment = main.state.orch.comment_preview(aid)

    assert "edited after this test plan" in comment


def test_jira_comment_does_not_flag_a_fresh_plan(client):
    import app.api.main as main

    aid = _import(client)
    client.post(f"/assessment/{aid}/design", follow_redirects=True)
    comment = main.state.orch.comment_preview(aid)

    assert "edited after this test plan" not in comment
